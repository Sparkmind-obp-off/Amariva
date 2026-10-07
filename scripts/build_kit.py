"""Build the real private paid artifact, embedded only in the server bundle.
Build-time Python only. Cloudflare runtime does not run this script.
Requires openpyxl. No customer data or third-party assets.
"""
from pathlib import Path
from io import BytesIO
import base64
import zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
ROOT = Path(__file__).resolve().parent.parent
wb = Workbook()
ws = wb.active
ws.title = 'Skenario'
rows = [
 ['AMARIVA — Pricing Decision Kit v1.0'],
 ['Isi sel hijau muda. Rumus dihitung saat workbook dibuka di Excel / LibreOffice.'],
 ['Input / hasil','Konservatif','Dasar','Optimistis'],
 ['Biaya variabel per unit',25000,25000,25000],
 ['Harga jual per unit',50000,50000,50000],
 ['Fee penjualan (%)',5,5,5],
 ['Biaya tetap per periode',2000000,2000000,2000000],
 ['Target unit per periode',80,200,300],
 [''],
 ['Fee per unit'],['Kontribusi per unit'],['Margin kontribusi'],['Titik impas (unit)'],['Estimasi laba operasional'],['Harga minimum untuk target unit'],
 [''],['Batasan: belum termasuk pajak, investasi awal, dan perubahan kapasitas.'],
 ['Kontribusi <= 0: BEP tidak tercapai. Harga minimum bukan rekomendasi harga pasar.']
]
for row in rows: ws.append(row)
for col in 'BCD':
 formulas = {10:f'={col}5*{col}6/100',11:f'={col}5-{col}4-{col}10',12:f'=IF({col}5>0,{col}11/{col}5,0)',13:f'=IF({col}11>0,ROUNDUP({col}7/{col}11,0),"Tidak tercapai")',14:f'={col}11*{col}8-{col}7',15:f'=IF(AND({col}8>0,{col}6<100),ROUNDUP(({col}4+{col}7/{col}8)/(1-{col}6/100),0),"Periksa input")'}
 for row,formula in formulas.items(): ws[f'{col}{row}']=formula
 for row in range(4,9): ws[f'{col}{row}'].fill=PatternFill('solid',fgColor='E5EED9')
 ws[f'{col}12'].number_format='0.00%'
 for row in [4,5,7,10,11,14,15]: ws[f'{col}{row}'].number_format='"Rp" #,##0;[Red]-"Rp" #,##0'
 dv=DataValidation(type='decimal',operator='between',formula1=0,formula2=999999999999);dv.error='Gunakan angka non-negatif.';dv.showErrorMessage=True;ws.add_data_validation(dv)
 for cell in [f'{col}4',f'{col}7']: dv.add(ws[cell])
 positive=DataValidation(type='decimal',operator='between',formula1=1,formula2=999999999999);positive.showErrorMessage=True;ws.add_data_validation(positive);positive.add(ws[f'{col}5'])
 fee=DataValidation(type='decimal',operator='between',formula1=0,formula2=99.99);fee.showErrorMessage=True;ws.add_data_validation(fee);fee.add(ws[f'{col}6'])
 units=DataValidation(type='whole',operator='between',formula1=1,formula2=10000000);units.showErrorMessage=True;ws.add_data_validation(units);units.add(ws[f'{col}8'])
ws.column_dimensions['A'].width=46
for col in 'BCD': ws.column_dimensions[col].width=24
ws.freeze_panes='B4'
ws.auto_filter.ref='A3:D15'
for cell in ws[3]: cell.fill=PatternFill('solid',fgColor='234537');cell.font=Font(color='FFFFFF',bold=True)
ws['A1'].font=Font(size=20,bold=True,color='234537')
check=wb.create_sheet('Audit biaya')
for row in [['Kategori','Item','Biaya','Periode / unit','Sumber data','Sudah dicek?'],['Variabel','Bahan',0,'per unit','',''],['Variabel','Kemasan',0,'per unit','',''],['Variabel','Produksi',0,'per unit','',''],['Variabel','Fee nominal pembayaran',0,'per unit','',''],['Tetap','Sewa',0,'per bulan','',''],['Tetap','Gaji tetap',0,'per bulan','',''],['Persentase','Fee platform',0,'persen harga jual','','']]:check.append(row)
experiment=wb.create_sheet('Log keputusan')
for row in [['Tanggal','Hipotesis','Harga','Unit target','Kontribusi target','Hasil aktual','Keputusan selanjutnya'],['','Contoh: harga baru menutup biaya tanpa menurunkan volume','','','','','']]: experiment.append(row)
for sheet in [check,experiment]:
 sheet.freeze_panes='A2'
 for cells in sheet.columns:sheet.column_dimensions[cells[0].column_letter].width=27
 for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='234537');cell.font=Font(color='FFFFFF',bold=True)
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
buf=BytesIO();wb.save(buf)
guide='''AMARIVA — Pricing Decision Kit v1.0

1. Buka pricing-workbook.xlsx di Excel / LibreOffice, atau impor ke Google Sheets.
2. Isi biaya variabel, harga, fee, biaya tetap, dan unit pada tiga skenario.
3. Biaya tetap dan target unit harus menggunakan periode yang sama.
4. Periksa tab Audit biaya. Masukkan fee nominal tetap ke biaya variabel.
5. Uji volume konservatif sebelum menetapkan harga. Gunakan Log keputusan.
6. Setelah penjualan nyata, bandingkan asumsi dan hasil, lalu perbarui harga.

Contoh default: harga 50.000, biaya 25.000, fee 5%, biaya tetap 2 juta.
Kontribusi per unit 22.500; BEP 89 unit.
80 unit: rugi 200.000. 200 unit: laba 2.500.000. 300 unit: laba 4.750.000.

Diskon 10% dari 50.000 ke 45.000 menurunkan kontribusi menjadi 17.750.
Untuk menyamai kontribusi awal pada 100 unit, perlu 127 unit.

Formula:
Kontribusi = harga - biaya variabel - harga * fee / 100.
BEP = ceil(biaya tetap / kontribusi), hanya jika kontribusi > 0.
Laba operasional estimasi = kontribusi * unit - biaya tetap.
Harga minimum = ceil((biaya variabel + biaya tetap / target unit) / (1-fee/100)).

Bukan kalkulator pajak, arus kas, atau modal kembali. Tidak ada jaminan laba.
Bantuan: buka /contact pada situs AMARIVA tempat Anda membeli.
'''
archive=BytesIO()
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
 for name,data in [('pricing-workbook.xlsx',buf.getvalue()),('START-HERE.txt',guide.encode()),('COST-CHECKLIST.txt','Cek bahan, kemasan, produksi, pengiriman, fee nominal, fee persen, sewa, gaji tetap, dan biaya langganan. Hindari menghitung satu biaya dua kali. Samakan periode biaya tetap dan target unit.'.encode())]:
  info=zipfile.ZipInfo(name,date_time=(2026,10,7,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;z.writestr(info,data)
(ROOT/'src/private-kit.ts').write_text('// Generated private server artifact. Never serve this module as a static asset.\nexport const kitBase64 = '+repr(base64.b64encode(archive.getvalue()).decode())+';\n')
print('Generated private kit:',len(archive.getvalue()),'bytes')
